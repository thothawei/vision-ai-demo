"""日報／週報：總檢驗數、良率（AI／最終）、NG 柏拉圖、各站別良率、複判一致率、配方檢驗、Cpk 摘要。
輸出 JSON／Excel（openpyxl）／可列印 HTML（瀏覽器列印成 PDF，不引入 PDF 套件）。"""

import html
import io
from datetime import date, datetime, timedelta

from core import recipes, spc, spc_service
from core.inspection_log import _final_verdict, get_stats, query_inspections


def period_range(period: str, day: date) -> tuple[str, str]:
    """day＝當天；week＝包含該日的那一週（週一到週日）。回傳 (date_from, date_to) 兩端含。"""
    if period == "day":
        return day.isoformat(), day.isoformat()
    if period == "week":
        monday = day - timedelta(days=day.weekday())
        return monday.isoformat(), (monday + timedelta(days=6)).isoformat()
    raise ValueError(f"period 只能是 day 或 week，目前是「{period}」")


def _yield(counts: dict) -> float | None:
    denom = counts["OK"] + counts["NG"]
    return round(counts["OK"] / denom, 4) if denom else None


def _by_station(records: list[dict]) -> list[dict]:
    groups: dict[str, dict] = {}
    for r in records:
        g = groups.setdefault(
            r["station"] or "（未填）", {"station": r["station"] or "（未填）", "total": 0, "OK": 0, "NG": 0, "INFO": 0}
        )
        g["total"] += 1
        g[_final_verdict(r)] += 1
    out = sorted(groups.values(), key=lambda g: g["station"])
    for g in out:
        g["yield"] = _yield(g)
    return out


def _cpk_summary(date_from: str, date_to: str) -> list[dict]:
    """期間內每個（料號，特性）的能力指標；有規格才算，資料不足就標示。"""
    out = []
    for c in spc_service.list_characteristics():
        rows = spc_service.load_series(c["part_no"], c["characteristic"], date_from, date_to)
        if not rows:
            continue
        spec = spc_service.current_spec(rows)
        entry = {
            "part_no": c["part_label"],
            "characteristic": c["characteristic"],
            "unit": c["unit"] or "",
            "count": len(rows),
            "lsl": spec["lsl"],
            "usl": spec["usl"],
            "cp": None,
            "cpk": None,
            "subgroups": 0,
            "reference_only": True,
            "note": "",
        }
        if spec["lsl"] is None and spec["usl"] is None:
            entry["note"] = "無規格"
        else:
            values = [r["value"] for r in rows]
            try:
                result = (
                    spc.xbar_r(values, 5, spec["lsl"], spec["usl"])
                    if len(values) >= 10
                    else spc.imr(values, spec["lsl"], spec["usl"])
                )
                cap = result["capability"]
                entry.update(
                    cp=cap["cp"],
                    cpk=cap["cpk"],
                    subgroups=result["subgroups"],
                    reference_only=cap["reference_only"],
                    note=cap["note"] or "",
                )
            except ValueError as e:
                entry["note"] = str(e)
        out.append(entry)
    return out


def build_report(period: str, day: date) -> dict:
    date_from, date_to = period_range(period, day)
    records = query_inspections(date_from=date_from, date_to=date_to, limit=1_000_000)
    stats = get_stats("module", date_from, date_to)
    defects = get_stats("defect", date_from, date_to)
    runs = recipes.get_run_stats("part", date_from, date_to)
    return {
        "period": period,
        "date_from": date_from,
        "date_to": date_to,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "summary": {
            "total": stats["total"],
            "ai_verdict_counts": stats["ai_verdict_counts"],
            "final_verdict_counts": stats["final_verdict_counts"],
            "ai_yield": stats["ai_yield"],
            "final_yield": stats["final_yield"],
            "reviewed_count": stats["reviewed_count"],
            "consistency_rate": stats["consistency_rate"],
            "adhoc_count": sum(1 for r in records if r.get("adhoc") == 1),
        },
        "by_module": stats["groups"],
        "pareto": defects["groups"],
        "by_station": _by_station(records),
        "runs": {
            "total": runs["total"],
            "ai_yield": runs["ai_yield"],
            "final_yield": runs["final_yield"],
            "final_verdict_counts": runs["final_verdict_counts"],
            "by_part": runs["groups"],
        },
        "cpk": _cpk_summary(date_from, date_to),
    }


# --- 輸出 ---


def _pct(v) -> str:
    return "—" if v is None else f"{v * 100:.1f}%"


def _num(v, digits=3) -> str:
    return "—" if v is None else f"{v:.{digits}f}"


def _title(report: dict) -> str:
    label = "日報" if report["period"] == "day" else "週報"
    span = report["date_from"] if report["period"] == "day" else f"{report['date_from']} ～ {report['date_to']}"
    return f"品質{label} {span}"


def to_xlsx(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "摘要"
    s = report["summary"]
    rows = [
        (_title(report),),
        ("產生時間", report["generated_at"]),
        (),
        ("總檢驗數", s["total"]),
        ("AI 良率", _pct(s["ai_yield"])),
        ("最終良率（含複判）", _pct(s["final_yield"])),
        ("AI 判定 OK／NG／INFO", "／".join(str(s["ai_verdict_counts"][k]) for k in ("OK", "NG", "INFO"))),
        ("最終判定 OK／NG／INFO", "／".join(str(s["final_verdict_counts"][k]) for k in ("OK", "NG", "INFO"))),
        ("已複判筆數", s["reviewed_count"]),
        ("AI／人工一致率", _pct(s["consistency_rate"])),
        ("臨時檢驗筆數（無工單或工單不在主檔）", s["adhoc_count"]),
        ("配方檢驗次數", report["runs"]["total"]),
        ("配方檢驗最終良率", _pct(report["runs"]["final_yield"])),
    ]
    for row in rows:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 28

    def sheet(name, header, data):
        sh = wb.create_sheet(name)
        sh.append(header)
        for cell in sh[1]:
            cell.font = Font(bold=True)
        for row in data:
            sh.append(row)
        for i in range(len(header)):
            sh.column_dimensions[chr(65 + i)].width = 22

    sheet(
        "NG 柏拉圖",
        ["NG 原因", "次數", "佔比 %", "累積 %"],
        [(g["label"], g["count"], g["pct"], g["cum_pct"]) for g in report["pareto"]],
    )
    sheet(
        "各站別良率",
        ["站別", "檢驗數", "OK", "NG", "INFO", "良率"],
        [(g["station"], g["total"], g["OK"], g["NG"], g["INFO"], _pct(g["yield"])) for g in report["by_station"]],
    )
    sheet(
        "各模組",
        ["模組", "檢驗數", "OK", "NG", "INFO", "良率"],
        [(g["module"], g["total"], g["OK"], g["NG"], g["INFO"], _pct(g["yield"])) for g in report["by_module"]],
    )
    sheet(
        "配方檢驗（依料號）",
        ["料號", "檢驗次數", "OK", "NG", "INFO", "良率"],
        [(g["part_no"], g["total"], g["OK"], g["NG"], g["INFO"], _pct(g["yield"])) for g in report["runs"]["by_part"]],
    )
    sheet(
        "Cpk 摘要",
        ["料號", "特性", "單位", "筆數", "規格下限", "規格上限", "Cp", "Cpk", "子群數", "備註"],
        [
            (
                c["part_no"],
                c["characteristic"],
                c["unit"],
                c["count"],
                c["lsl"],
                c["usl"],
                _num(c["cp"]),
                _num(c["cpk"]),
                c["subgroups"],
                ("僅供參考（子群少於 25）" if c["reference_only"] and c["cpk"] is not None else "")
                + (" " + c["note"] if c["note"] else ""),
            )
            for c in report["cpk"]
        ],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_html(report: dict) -> str:
    e = html.escape

    def table(header, rows):
        head = "".join(f"<th>{e(h)}</th>" for h in header)
        body = "".join("<tr>" + "".join(f"<td>{e(str(c))}</td>" for c in row) + "</tr>" for row in rows)
        return f"<table><thead><tr>{head}</tr></thead><tbody>{body or '<tr><td colspan=99>無資料</td></tr>'}</tbody></table>"

    s = report["summary"]
    cards = "".join(
        f"<div class='card'><div class='label'>{e(k)}</div><div class='value'>{e(str(v))}</div></div>"
        for k, v in (
            ("總檢驗數", s["total"]),
            ("AI 良率", _pct(s["ai_yield"])),
            ("最終良率", _pct(s["final_yield"])),
            ("已複判", s["reviewed_count"]),
            ("AI／人工一致率", _pct(s["consistency_rate"])),
            ("臨時檢驗", s["adhoc_count"]),
            ("配方檢驗次數", report["runs"]["total"]),
        )
    )
    sections = [
        (
            "NG 原因柏拉圖",
            ["NG 原因", "次數", "佔比 %", "累積 %"],
            [(g["label"], g["count"], g["pct"], g["cum_pct"]) for g in report["pareto"]],
        ),
        (
            "各站別良率",
            ["站別", "檢驗數", "OK", "NG", "INFO", "良率"],
            [(g["station"], g["total"], g["OK"], g["NG"], g["INFO"], _pct(g["yield"])) for g in report["by_station"]],
        ),
        (
            "各模組",
            ["模組", "檢驗數", "OK", "NG", "INFO", "良率"],
            [(g["module"], g["total"], g["OK"], g["NG"], g["INFO"], _pct(g["yield"])) for g in report["by_module"]],
        ),
        (
            "Cpk 摘要",
            ["料號", "特性", "筆數", "規格", "Cp", "Cpk", "備註"],
            [
                (
                    c["part_no"],
                    c["characteristic"],
                    c["count"],
                    f"{_num(c['lsl'], 2)} ～ {_num(c['usl'], 2)}",
                    _num(c["cp"]),
                    _num(c["cpk"]),
                    ("僅供參考（子群少於 25）；" if c["reference_only"] and c["cpk"] is not None else "") + c["note"],
                )
                for c in report["cpk"]
            ],
        ),
    ]
    body = "".join(f"<h2>{e(t)}</h2>{table(h, r)}" for t, h, r in sections)
    return f"""<!DOCTYPE html><html lang="zh-Hant"><head><meta charset="UTF-8"><title>{e(_title(report))}</title>
<style>body{{font-family:-apple-system,"PingFang TC","Microsoft JhengHei",sans-serif;margin:24px;color:#1d1d1f}}
h1{{font-size:22px}} h2{{font-size:16px;margin-top:24px}} .cards{{display:flex;flex-wrap:wrap;gap:12px}}
.card{{border:1px solid #ccc;border-radius:8px;padding:8px 14px;min-width:110px}} .label{{font-size:12px;color:#666}}
.value{{font-size:20px;font-weight:600}} table{{border-collapse:collapse;width:100%;font-size:13px}}
th,td{{border:1px solid #ccc;padding:5px 8px;text-align:left}} th{{background:#f0f0f0}}
.meta{{color:#666;font-size:12px}} @media print{{body{{margin:8mm}} .noprint{{display:none}}}}</style></head><body>
<h1>{e(_title(report))}</h1><div class="meta">產生時間 {e(report["generated_at"])}　｜　<span class="noprint">列印成 PDF：瀏覽器「列印」→「另存為 PDF」</span></div>
<div class="cards">{cards}</div>{body}</body></html>"""
