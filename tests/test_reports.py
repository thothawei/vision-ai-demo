"""Phase 20：日報／週報（JSON、Excel、可列印 HTML）、排程腳本、歷史量測值回填。"""

import io
import json
import sys
from datetime import date, timedelta
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from core import reports
from core.inspection_log import connect, log_inspection, update_review

ROOT = Path(__file__).resolve().parent.parent
TODAY = date.today()


def _seed():
    """今天：站別 S1 兩筆 OK 一筆 NG（NG 被複判成 OK）、站別 S2 一筆 NG；一筆沒有站別的臨時檢驗；量測值 15 筆。"""
    ids = {}
    ids["ok1"] = log_inspection("codes", "OK", "e", 1, [{"效期狀態": "正常"}], {"station": "S1", "work_order": None})
    ids["ok2"] = log_inspection("codes", "OK", "e", 1, [{"效期狀態": "正常"}], {"station": "S1"})
    ids["ng_reviewed"] = log_inspection("codes", "NG", "e", 1, [{"效期狀態": "已過期"}], {"station": "S1"})
    ids["ng"] = log_inspection("codes", "NG", "e", 1, [{"效期狀態": "已過期"}], {"station": "S2"})
    ids["nostation"] = log_inspection("general", "INFO", "e", 1, [{"主要物件": "x"}], {})
    update_review(ids["ng_reviewed"], "OK", "qc", "誤判")
    for v in [9, 10, 11, 12, 13, 10, 11, 12, 13, 14, 8, 9, 10, 11, 12]:
        log_inspection("measure", "OK", "e", 1, [{"長度mm": float(v), "目標長度mm": 10.5, "公差mm": 5.5}],
                       {"part_no": "P-1", "station": "S1"})
    return ids


def test_period_ranges():
    assert reports.period_range("day", date(2026, 9, 29)) == ("2026-09-29", "2026-09-29")
    assert reports.period_range("week", date(2026, 9, 29)) == ("2026-09-28", "2026-10-04")  # 週二 → 週一到週日
    assert reports.period_range("week", date(2026, 9, 28)) == ("2026-09-28", "2026-10-04")
    assert reports.period_range("week", date(2026, 10, 4)) == ("2026-09-28", "2026-10-04")
    with pytest.raises(ValueError):
        reports.period_range("month", date(2026, 9, 29))


def test_report_numbers_are_correct(client):
    _seed()
    r = reports.build_report("day", TODAY)
    s = r["summary"]
    assert s["total"] == 20
    assert s["ai_verdict_counts"] == {"OK": 17, "NG": 2, "INFO": 1}
    assert s["final_verdict_counts"] == {"OK": 18, "NG": 1, "INFO": 1}  # 一筆 NG 複判成 OK
    assert s["ai_yield"] == pytest.approx(17 / 19, abs=1e-4) and s["final_yield"] == pytest.approx(18 / 19, abs=1e-4)
    assert s["reviewed_count"] == 1 and s["consistency_rate"] == 0.0  # 複判結果與 AI 不一致
    assert s["adhoc_count"] == 20  # 全都沒有主檔工單
    stations = {g["station"]: g for g in r["by_station"]}
    assert stations["S1"]["total"] == 18 and stations["S1"]["NG"] == 0 and stations["S1"]["yield"] == 1.0  # 最終判定
    assert stations["S2"]["NG"] == 1 and stations["S2"]["yield"] == 0.0
    assert stations["（未填）"]["INFO"] == 1 and stations["（未填）"]["yield"] is None
    assert [(g["label"], g["count"]) for g in r["pareto"]] == [("追溯碼已過期", 2)]
    modules = {g["module"]: g for g in r["by_module"]}
    assert modules["codes"]["total"] == 4 and modules["measure"]["total"] == 15
    cpk = r["cpk"][0]
    assert (cpk["part_no"], cpk["characteristic"], cpk["count"], cpk["subgroups"]) == ("P-1", "長度", 15, 3)
    assert cpk["cp"] == pytest.approx(1.06608, rel=1e-4) and cpk["cpk"] == pytest.approx(0.96917, rel=1e-4)
    assert cpk["reference_only"] is True


def test_report_only_covers_the_requested_period(client):
    _seed()
    yesterday = reports.build_report("day", TODAY - timedelta(days=1))
    assert yesterday["summary"]["total"] == 0 and yesterday["cpk"] == [] and yesterday["pareto"] == []
    with connect() as conn:  # 把三筆挪到 8 天前：當天報表少三筆，但仍在「包含它的那一週」報表內
        old = (TODAY - timedelta(days=8)).isoformat() + "T10:00:00"
        conn.execute("UPDATE inspections SET created_at = ? WHERE id IN (1, 2, 3)", (old,))
        conn.commit()
    assert reports.build_report("day", TODAY)["summary"]["total"] == 17
    assert reports.build_report("day", TODAY - timedelta(days=8))["summary"]["total"] == 3
    week = reports.build_report("week", TODAY - timedelta(days=8))
    assert week["summary"]["total"] >= 3 and week["date_from"] <= (TODAY - timedelta(days=8)).isoformat() <= week["date_to"]


def test_empty_report_does_not_crash(client):
    r = reports.build_report("day", TODAY)
    assert r["summary"]["total"] == 0 and r["summary"]["ai_yield"] is None
    assert reports.to_xlsx(r) and "無資料" in reports.to_html(r)


def test_excel_is_readable_and_keeps_chinese(client):
    _seed()
    wb = load_workbook(io.BytesIO(reports.to_xlsx(reports.build_report("day", TODAY))))
    assert wb.sheetnames == ["摘要", "NG 柏拉圖", "各站別良率", "各模組", "配方檢驗（依料號）", "Cpk 摘要"]
    summary = {row[0].value: row[1].value for row in wb["摘要"].iter_rows(min_row=4) if row[0].value}
    assert summary["總檢驗數"] == 20 and summary["AI 良率"] == "89.5%" and summary["最終良率（含複判）"] == "94.7%"
    assert wb["摘要"]["A1"].value.startswith("品質日報")
    assert [c.value for c in wb["NG 柏拉圖"][2]] == ["追溯碼已過期", 2, 100.0, 100.0]
    station_rows = {row[0].value: [c.value for c in row] for row in wb["各站別良率"].iter_rows(min_row=2)}
    assert station_rows["S2"][:5] == ["S2", 1, 0, 1, 0]
    cpk_row = [c.value for c in wb["Cpk 摘要"][2]]
    assert cpk_row[:4] == ["P-1", "長度", "mm", 15] and cpk_row[7] == "0.969" and "僅供參考" in cpk_row[9]


def test_html_is_printable_and_escapes_content(client):
    log_inspection("defect", "NG", "e", 1,
                   [{"明細": [{"瑕疵類型": "<script>alert(1)</script>"}]}], {"station": "<b>S</b>"})
    page = reports.to_html(reports.build_report("day", TODAY))
    assert "<script>alert(1)</script>" not in page and "&lt;script&gt;" in page and "&lt;b&gt;S&lt;/b&gt;" in page
    assert "@media print" in page and "<title>品質日報" in page


def test_report_api_formats_periods_and_validation(admin_client):
    _seed()
    j = admin_client.get("/api/reports/summary", params={"date": TODAY.isoformat()})
    assert j.status_code == 200 and j.json()["summary"]["total"] == 20 and j.json()["period"] == "day"
    week = admin_client.get("/api/reports/summary", params={"period": "week", "date": "2026-09-29"}).json()
    assert (week["date_from"], week["date_to"]) == ("2026-09-28", "2026-10-04")
    x = admin_client.get("/api/reports/summary", params={"format": "xlsx"})
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    assert f"report-day-{TODAY.isoformat()}.xlsx" in x.headers["content-disposition"]
    assert load_workbook(io.BytesIO(x.content)).sheetnames[0] == "摘要"
    h = admin_client.get("/api/reports/summary", params={"format": "html"})
    assert h.status_code == 200 and h.headers["content-type"].startswith("text/html") and "品質日報" in h.text
    assert admin_client.get("/api/reports/summary", params={"format": "pdf"}).status_code == 400
    assert admin_client.get("/api/reports/summary", params={"period": "year"}).status_code == 400
    assert admin_client.get("/api/reports/summary", params={"date": "not-a-date"}).status_code == 422


def test_report_permissions(admin_client, qc_client, operator_client, client):
    assert qc_client.get("/api/reports/summary").status_code == 200
    assert client.get("/api/reports/summary").status_code == 200  # API key 唯讀
    assert operator_client.get("/api/reports/summary").status_code == 403
    assert TestClient(admin_client.app).get("/api/reports/summary").status_code == 401


# --- 腳本 ---

def test_daily_report_script_writes_excel_and_html(client, tmp_path):
    _seed()
    sys.path.insert(0, str(ROOT / "scripts"))
    import daily_report

    paths = daily_report.generate("day", TODAY, tmp_path / "out")
    assert [p.suffix for p in paths] == [".xlsx", ".html"] and all(p.stat().st_size > 0 for p in paths)
    assert load_workbook(paths[0]).sheetnames[0] == "摘要"
    assert "品質日報" in paths[1].read_text(encoding="utf-8")


def test_reports_directory_is_gitignored():
    import subprocess

    def ignored(path):
        return subprocess.run(["git", "check-ignore", "-q", path], cwd=ROOT).returncode == 0

    assert ignored("reports/daily.xlsx")
    # 規則要錨定在根目錄：曾因 `reports/`、`models/` 沒加 `/`，把 backend/modules 下的同名原始碼目錄也擋掉而漏 commit
    assert not ignored("backend/modules/reports/router.py")
    assert not ignored("backend/modules/models/router.py")


def _legacy_rows():
    """模擬 Phase 20 之前的紀錄：直接寫 inspections，不會產生 measurements。"""
    with connect() as conn:
        for module, summary, part in (
            ("measure", [{"長度mm": 50.6, "寬度mm": 20.3, "孔徑mm": [5.4], "目標長度mm": 50, "長度誤差mm": 0.6}], "P-9"),
            ("colordiff", [{"ΔE00": 1.7, "門檻": 3.0, "判定": "合格"}], "P-9"),
            ("measure", [{"計數結果": 8, "預期數量": 8, "明細": []}], "P-9"),
            ("codes", [{"效期狀態": "正常"}], "P-9"),
        ):
            conn.execute("INSERT INTO inspections (created_at, module, verdict, engine, elapsed_ms, summary_json, part_no)"
                         " VALUES ('2026-01-05T09:00:00', ?, 'OK', 'e', 1, ?, ?)", (module, json.dumps(summary, ensure_ascii=False), part))
        conn.commit()


def test_backfill_dry_run_writes_nothing_then_real_run_is_idempotent(client):
    sys.path.insert(0, str(ROOT / "scripts"))
    import backfill_measurements

    _legacy_rows()
    dry = backfill_measurements.backfill(dry_run=True)
    assert dry == {"inspections": 2, "measurements": 4, "dry_run": True}  # 長度、寬度、孔徑＋ΔE
    with connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM measurements").fetchone()[0] == 0

    done = backfill_measurements.backfill()
    assert (done["inspections"], done["measurements"]) == (2, 4)
    with connect() as conn:
        rows = {r["characteristic"]: dict(r) for r in conn.execute("SELECT * FROM measurements").fetchall()}
    assert set(rows) == {"長度", "寬度", "孔徑", "色差ΔE00"}
    assert rows["長度"]["lsl"] is None and rows["長度"]["usl"] is None and rows["長度"]["target"] == 50  # 當時沒存公差
    assert rows["色差ΔE00"]["usl"] == 3.0 and rows["長度"]["created_at"] == "2026-01-05T09:00:00"
    assert rows["長度"]["part_no"] == "P-9"
    assert backfill_measurements.backfill() == {"inspections": 0, "measurements": 0, "dry_run": False}  # 再跑不會重複
